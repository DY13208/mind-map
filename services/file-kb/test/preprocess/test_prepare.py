import hashlib
import json
import importlib.util
import tempfile
import unittest
from pathlib import Path
from PIL import Image
import openpyxl
import xlwt

SPEC = importlib.util.spec_from_file_location("prepare", Path(__file__).parents[2] / "src" / "prepare.py")
prepare = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(prepare)

class PrepareTests(unittest.TestCase):
    def test_image_dpi_and_original(self):
        with tempfile.TemporaryDirectory() as folder:
            for dpi in (1200, 300, None):
                src, dst = Path(folder)/"source.jpg", Path(folder)/"output.png"
                Image.new("RGB", (800, 584), "white").save(src, **({"dpi": (dpi,dpi)} if dpi else {}))
                before = hashlib.sha256(src.read_bytes()).digest()
                stats = prepare.image(src, dst)
                with Image.open(dst) as im:
                    self.assertEqual(im.size, (800,584))
                    self.assertEqual(im.mode, "RGB")
                    self.assertAlmostEqual(im.info['dpi'][0], 216, delta=.02)
                self.assertEqual(stats['width'], 800)
                self.assertEqual(before, hashlib.sha256(src.read_bytes()).digest())

    def test_exif_rational_dpi_stats_are_json_serializable(self):
        with tempfile.TemporaryDirectory() as folder:
            src, dst = Path(folder)/'exif.jpg', Path(folder)/'out.png'
            from PIL.TiffImagePlugin import IFDRational
            exif=Image.Exif();exif[282]=IFDRational(1200,1);exif[283]=IFDRational(1200,1);exif[296]=2
            Image.new('RGB',(80,60),'white').save(src,exif=exif)
            self.assertEqual(json.loads(json.dumps(prepare.image(src,dst)))['originalDpi'],[1200,1200])

    def test_orientation_transparency_and_limits(self):
        with tempfile.TemporaryDirectory() as folder:
            src, dst = Path(folder)/"source.png", Path(folder)/"out.png"
            exif = Image.Exif(); exif[274] = 6
            Image.new("RGBA", (20,10), (0,0,0,0)).save(src, exif=exif)
            prepare.image(src,dst)
            with Image.open(dst) as im:
                self.assertEqual(im.size,(10,20)); self.assertEqual(im.getpixel((0,0)),(255,255,255))
            with self.assertRaisesRegex(ValueError, "像素"):
                prepare.image(src,dst,max_pixels=100)
            src.write_bytes(b"broken")
            with self.assertRaises(Exception): prepare.image(src,dst)

    def test_xls_types_dates_merges_and_literal_formula(self):
        for dates_1904 in (False, True):
            with tempfile.TemporaryDirectory() as folder:
                src,dst=Path(folder)/'in.xls',Path(folder)/'out.xlsx'
                wb=xlwt.Workbook();wb.dates_1904=dates_1904
                sh=wb.add_sheet('中文明细');sh.write_merge(0,0,0,2,'合并标题')
                sh.write(2,0,'=literal');sh.write(2,1,False);sh.write(2,2,12.5)
                import datetime
                sh.write(5,0,datetime.datetime(2024,5,6),xlwt.easyxf(num_format_str='YYYY-MM-DD'))
                sh.write(5,1,xlwt.Formula('1+2'))
                wb.add_sheet('第二页').write(0,0,'内容');wb.save(str(src))
                stats=prepare.xls(src,dst)
                converted=openpyxl.load_workbook(dst,data_only=False)
                self.assertEqual(converted.sheetnames,['中文明细','第二页'])
                s=converted.worksheets[0]
                self.assertEqual(str(next(iter(s.merged_cells.ranges))),'A1:C1')
                self.assertEqual(s['A3'].value,'=literal');self.assertEqual(s['A3'].data_type,'s')
                self.assertIs(s['B3'].value,False);self.assertEqual(s['C3'].value,12.5)
                self.assertEqual(s['A6'].value,datetime.datetime(2024,5,6))
                self.assertEqual(stats['sheets'],2)
                self.assertEqual(s['B6'].value,None) # xlwt provides no cached formula result
                converted.close()

    def test_corrupt_xls_is_rejected_and_no_output_is_published(self):
        with tempfile.TemporaryDirectory() as folder:
            src,dst=Path(folder)/'in.xls',Path(folder)/'out.xlsx'
            src.write_bytes(b'not a binary workbook')
            with self.assertRaisesRegex(ValueError,'二进制'):
                prepare.xls(src,dst)
            self.assertFalse(dst.exists())
            src.write_bytes(bytes.fromhex('d0cf11e0a1b11ae1')+b'corrupt')
            with self.assertRaises(Exception): prepare.xls(src,dst)
            self.assertFalse(dst.exists())

if __name__ == '__main__': unittest.main()
