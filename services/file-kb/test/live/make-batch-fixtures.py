"""Synthetic, non-business fixtures. Never reads an existing business file."""
import os
from pathlib import Path
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation
from openpyxl import Workbook

lab = Path(os.environ['KB_LAB_ROOT']).resolve()
if not str(lab).startswith('D:\\File-KB-Lab-'):
    raise RuntimeError('dedicated lab directory required')
root = lab / 'fixtures'
root.mkdir(exist_ok=True)
for i in range(100):
    ext = ['pdf', 'docx', 'xlsx', 'pptx', 'md'][i % 5]
    p = root / f'lab-{i:03}.{ext}'
    text = f'File KB synthetic acceptance {i}. Inventory SKU-{i:03}, price 12.50, stock 25. This is generated non-business data.'
    if ext == 'pdf':
        c = canvas.Canvas(str(p))
        c.drawString(40, 800, text[:105])
        c.drawString(40, 770, 'SKU | Price | Stock | Date')
        c.drawString(40, 750, f'SKU-{i:03} | 12.50 | 25 | 2026-10-19')
        c.showPage()
        c.drawString(40, 800, f'Page two acceptance {i}')
        c.save()
    elif ext == 'docx':
        d = Document()
        d.add_heading(f'Inventory {i}', 0)
        d.add_paragraph(text)
        t = d.add_table(rows=1, cols=3)
        for cell, value in zip(t.rows[0].cells, ['SKU', 'Price', 'Stock']):
            cell.text = value
        for cell, value in zip(t.add_row().cells, [f'SKU-{i:03}', '12.50', '25']):
            cell.text = value
        d.save(p)
    elif ext == 'xlsx':
        w = Workbook()
        s = w.active
        s.title = 'Inventory'
        s.append(['SKU', 'Description', 'Price', 'Stock', 'Date'])
        for j in range(200):
            s.append([f'SKU-{i:03}-{j}', text, 12.5+j, 25, '2026-10-19'])
        w.save(p)
    elif ext == 'pptx':
        r = Presentation()
        s = r.slides.add_slide(r.slide_layouts[1])
        s.shapes.title.text = f'Inventory {i}'
        s.placeholders[1].text = text
        r.save(p)
    else:
        p.write_text(f'# Inventory {i}\n\n{text}\n\n| SKU | Price | Stock |\n|---|---|---|\n| SKU-{i:03} |12.50|25|', encoding='utf8')
print('fixtures:', len(list(root.iterdir())))
