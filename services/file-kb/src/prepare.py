"""Parse-only normalization. Never modifies the uploaded source file."""
import json
import sys
from pathlib import Path


def image(source, destination, max_pixels=40_000_000):
    from PIL import Image, ImageOps
    with Image.open(source) as original:
        width, height = original.size
        if width * height > max_pixels:
            raise ValueError("图片超过 4000 万像素，无法安全解析")
        dpi = original.info.get('dpi')
        oriented = ImageOps.exif_transpose(original)
        rgba = oriented.convert('RGBA')
        white = Image.new('RGBA', rgba.size, 'white')
        white.alpha_composite(rgba)
        normalized = white.convert('RGB')
        # Strip source EXIF/ICC as these may override the normalized DPI or orientation.
        normalized.info.clear()
        normalized.save(destination, format='PNG', dpi=(216, 216))
        return {'method': 'image-normalize', 'originalWidth': width, 'originalHeight': height,
                'originalDpi': [float(value) for value in dpi] if dpi else None, 'width': normalized.width,
                'height': normalized.height, 'dpi': 216}


def xls(source, destination):
    import xlrd
    import openpyxl
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
    from openpyxl.utils.exceptions import IllegalCharacterError
    with open(source, 'rb') as stream:
        if stream.read(8) != bytes.fromhex('d0cf11e0a1b11ae1'):
            raise ValueError('XLS 格式无效：需要二进制 Excel 工作簿')
    book = xlrd.open_workbook(str(source), formatting_info=True, on_demand=True)
    output = openpyxl.Workbook()
    output.remove(output.active)
    output.epoch = CALENDAR_MAC_1904 if book.datemode else CALENDAR_WINDOWS_1900
    try:
        for source_sheet in book.sheets():
            sheet = output.create_sheet(source_sheet.name)
            for row in range(source_sheet.nrows):
                for col in range(source_sheet.ncols):
                    cell = source_sheet.cell(row, col)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#VALUE!')
                    target = sheet.cell(row+1, col+1)
                    try:
                        target.value = value
                    except IllegalCharacterError as exc:
                        raise ValueError('XLS 包含无法转换的控制字符') from exc
                    if cell.ctype == xlrd.XL_CELL_TEXT:
                        target.data_type = 's'  # Literal '=...' must never become an executable formula.
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        target.data_type = 'e'
            for top, bottom, left, right in source_sheet.merged_cells:
                sheet.merge_cells(start_row=top+1, end_row=bottom, start_column=left+1, end_column=right)
        output.save(destination)
        return {'method': 'xls-to-xlsx', 'sheets': book.nsheets, 'dateSystem': 1904 if book.datemode else 1900}
    finally:
        output.close()
        book.release_resources()


def main():
    try:
        mode, source, destination = sys.argv[1:]
        if mode not in ('image', 'xls'):
            raise ValueError('未知预处理类型')
        print(json.dumps(globals()[mode](Path(source), Path(destination)), ensure_ascii=False))
    except Exception as exc:
        # Do not include parser tracebacks, cell contents or uploaded filenames in diagnostics.
        detail = str(exc) if isinstance(exc, ValueError) else '无法读取、文件损坏、已加密或格式不受支持'
        print(json.dumps({'error': '文件预处理失败：' + detail}), file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
