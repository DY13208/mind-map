#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate D：天猫运营 execution sheet (HTML) + monitoring workbook (XLSX)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TS = "2026-09-14_1708"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "output")
OUT = os.path.abspath(OUT)
os.makedirs(OUT, exist_ok=True)

HTML_NAME = f"D_天猫运营_执行单_{TS}.html"
XLSX_NAME = f"D_天猫运营_监控台账_{TS}.xlsx"

SOFTP = "2a6d0264-ab20-41bd-ae44-f70279be9ae1"
TRACE = "tr-1a08ebb1f55-69797395"

# ---- SOP content (verbatim from map subtree) ----
activity_types = ["平台大促", "百亿补贴", "淘金币", "U先试用"]
activity_rows = [
    ("确认平台活动节奏排期", "步骤", "对齐平台大促/活动日历，锁定报名窗口", "平台活动日历", "待执行"),
    ("报名活动", "步骤", "报名成功 / 拿到活动回执", "天猫商家中心", "待执行"),
    ("活动主图和赠品权益设置", "步骤", "主图与赠品权益按活动要求配置", "店铺后台", "待执行"),
    ("客服话术确认", "步骤", "活动话术对齐并下发客服", "客服组", "待执行"),
    ("检查定时设置", "步骤", "开售/结束定时、上下架时间正确", "店铺后台", "待执行"),
    ("活动上线", "步骤", "活动已按排期上线", "店铺后台", "待执行"),
    ("活动检查 · 前端画面", "检查", "前端展示正确（无错图/错价/错权益）", "前端自检", "待检查"),
    ("活动检查 · 价格设置", "检查", "活动价与优惠券/满减生效正确", "生意参谋/后台", "待检查"),
    ("活动检查 · 赠品设置", "检查", "赠品规则生效、库存充足", "店铺后台", "待检查"),
    ("活动检查 · 监控库存，防止超卖和下架", "检查", "库存水位监控，防超卖/防意外下架", "吉客云/库存", "待检查"),
    ("活动结束和前端恢复检查", "步骤", "活动结束后前端恢复日常，无残留活动价", "店铺后台", "待执行"),
]

huopin_rows = [
    ("对应品牌战略的产品定位", "产品定位", "与品牌战略一致的产品定位", "品牌战略", "待确认"),
    ("核心产品链接新增SKU，提高销售承接能力", "链接设置", "核心链接新增 SKU 已生效", "店铺后台", "待执行"),
    ("核心产品备份多链接，承接多活动/明星/达播落地", "链接设置", "备份链接就绪可承接流量", "店铺后台", "待执行"),
    ("新品链接规划", "链接设置", "新品链接规划完成", "店铺后台", "待执行"),
    ("常规和滞销产品链接保持，在热销链接并入常规和滞销产品，提高连带率", "链接设置", "热销链接已并入常规/滞销品，连带率提升", "生意参谋", "待执行"),
    ("链接是否成功验证 · 链接转化率提升", "验证", "转化率环比提升", "生意参谋", "待取数"),
    ("链接是否成功验证 · 链接客单价提升", "验证", "客单价环比提升", "生意参谋", "待取数"),
    ("链接是否成功验证 · 链接总销售额提升", "验证", "总销售额环比提升", "生意参谋", "待取数"),
    ("优化卖点", "优化页面", "卖点已优化上线", "店铺后台", "待执行"),
    ("检查文案，符合大贸备案要求", "优化页面", "文案符合大贸备案要求（外链表核对）", "大贸备案表", "待取数"),
    ("页面是否合规和满足消费者需求的验证 · 符合中国大陆法规要求", "验证", "符合中国大陆法规", "法务/备案", "待检查"),
    ("页面是否合规和满足消费者需求的验证 · 无打假漏洞", "验证", "无打假风险点", "法务/备案", "待检查"),
    ("页面是否合规和满足消费者需求的验证 · 商品转化率提升", "验证", "商品转化率提升", "生意参谋", "待取数"),
]

zhibo_rows = [
    ("达播 · 商务BD", "达播流程", "达人筛选/接洽完成", "商务BD", "待执行"),
    ("达播 · 提交合同和OFFER", "达播流程", "合同与 OFFER 提交完成", "商务BD", "待执行"),
    ("达播 · 确认排期和上播", "达播流程", "排期确认、按约上播", "商务BD", "待执行"),
    ("达播 · 上播跟进", "达播流程", "上播过程跟进记录", "直播组", "待执行"),
    ("达播 · 下播后复盘", "达播流程", "下播复盘完成", "直播组", "待执行"),
    ("达播成果验证 · ROI=3 完成 → 沟通长期合作", "验证", "达播 ROI ≥ 3，转长期合作", "达播结算后台", "待取数"),
    ("达播成果验证 · ROI=3 未完成 → 退款 / 继续补播&挂链销售", "验证", "达播 ROI < 3，执行退款或补播&挂链", "达播结算后台", "待取数"),
    ("店播 · 店播规划", "店播流程", "店播规划完成", "直播组", "待执行"),
    ("店播 · 店播切片素材运营", "店播流程", "切片素材持续运营", "直播组", "待执行"),
    ("店播 · 店播中控和付费推广", "店播流程", "中控与付费推广执行", "直播组", "待执行"),
    ("店播成果验证 · 定播ROI=3 完成 → 进一步优化店播销售表现", "验证", "定播 ROI ≥ 3，继续优化", "店播中控后台", "待取数"),
    ("店播成果验证 · 定播ROI=3 未完成 → 店播提升计划", "验证", "定播 ROI < 3，出店播提升计划", "店播中控后台", "待取数"),
]

huiyuan_rows = [
    ("新会员招揽提升规划", "新会员规划", "招揽提升规划完成", "会员中心", "待执行"),
    ("新会员权益", "新会员规划", "新会员权益配置完成", "会员中心", "待执行"),
    ("新会员规划验证 · 新会员入会率", "验证", "新会员入会率达标", "会员中心", "待取数"),
    ("新会员规划验证 · 新会员成交占比", "验证", "新会员成交占比达标", "会员中心", "待取数"),
    ("会员登记和兑换权益", "会员复购", "会员登记与权益兑换配置", "会员中心", "待执行"),
    ("会员营销活动", "会员复购", "会员营销活动落地", "会员中心", "待执行"),
    ("会员成交验证 · 会员复购率", "验证", "会员复购率达标", "会员中心", "待取数"),
    ("会员成交验证 · 会员整体销售占比", "验证", "会员整体销售占比达标", "会员中心", "待取数"),
]

criteria = [
    ("达播成果验证", "达播 ROI = 3", "ROI ≥ 3 → 沟通长期合作", "ROI < 3 → 退款 / 继续补播&挂链销售", "达播结算后台（导图未挂）"),
    ("店播成果验证", "定播 ROI = 3", "ROI ≥ 3 → 进一步优化店播销售表现", "ROI < 3 → 店播提升计划", "店播中控后台（导图未挂）"),
    ("链接是否成功验证", "转化率 / 客单价 / 总销售额 三项提升", "视为链接成功", "未提升 → 复盘调整链接结构", "生意参谋（导图未挂）"),
    ("页面合规验证", "符合中国大陆法规 / 无打假漏洞 / 商品转化率提升", "合规通过可上线", "不合规 → 整改页面", "大贸备案表（登录墙）"),
    ("会员验证", "新会员入会率/成交占比；会员复购率/整体销售占比", "指标达成", "未达成 → 加强会员复购运营", "天猫会员中心（导图未挂）"),
]

modules = [
    ("D：活动运营", "b58e3476", 18, "平台大促 / 百亿补贴 / 淘金币 / U先试用"),
    ("D：货品运营", "b5c3d81f", 18, "产品定位 / 链接设置 / 优化页面"),
    ("D：直播运营", "d9f7a30a", 25, "达播 / 店播"),
    ("D：会员运营", "e8d2e472", 12, "新会员规划 / 会员复购"),
]

actions = [
    ("取数", "登录天猫生意参谋 + 千牛，导出个护部天猫店（13CHAN10030）近 30 天链接级 转化率 / 客单价 / GMV，回填台账「货品运营」页", "周中昱"),
    ("取数", "汇总达播 / 店播结算数据，核算 达播 ROI、定播 ROI 是否达到 3，回填台账「直播运营」页", "直播组"),
    ("取数", "天猫会员中心导出 新会员入会率 / 成交占比、会员复购率 / 整体销售占比，回填台账「会员运营」页", "会员运营"),
    ("口径", "在已登录企业微信中打开「大贸备案表」核对文案/备案口径，导出或回贴本单", "周中昱"),
    ("活动", "按平台节奏锁定双11/大促报名窗口，走查「活动步骤」7 步；上线后按「活动检查」4 项（前端画面 / 价格 / 赠品 / 库存防超卖）逐项打勾", "周中昱"),
    ("验证", "每项按「判定口径表」判定 完成/未完成，未完成进入对应分支动作（补播·退款 / 店播提升计划 / 整改页面）", "周中昱"),
]

# ---------------------------------------------------------------- HTML
def pill(status):
    if status in ("待取数",):
        return '<span class="pill p-todo">待取数</span>'
    if status in ("待检查", "待确认"):
        return '<span class="pill p-info">%s</span>' % status
    if status == "待执行":
        return '<span class="pill p-core">待执行</span>'
    return '<span class="pill">%s</span>' % status

def rows_html(rows, module):
    out = []
    first = True
    for i, (name, typ, judge, src, st) in enumerate(rows, 1):
        mod_cell = f'<td rowspan="{len(rows)}" class="modcell">{module}</td>' if first else ""
        out.append(
            f'<tr>{mod_cell}<td class="num">{i}</td><td>{name}</td>'
            f'<td><span class="pill p-info">{typ}</span></td>'
            f'<td>{judge}</td><td class="src">{src}</td><td>{pill(st)}</td></tr>'
        )
        first = False
    return "\n".join(out)

mod_rows = modules_row_html = "\n".join(
    f'<tr><td>{n}</td><td class="src" style="font-family:ui-monospace,Consolas,monospace">{u}</td>'
    f'<td class="num">{c}</td><td>{d}</td></tr>'
    for n, u, c, d in modules
)

criteria_rows = "\n".join(
    f'<tr><td><b>{a}</b></td><td>{b}</td><td class="ok">{c}</td><td class="bad">{d}</td><td class="src">{e}</td></tr>'
    for a, b, c, d, e in criteria
)

actions_html = "\n".join(
    f'<li><b>{a}</b> — {b}<span class="who">{c}</span></li>' for a, b, c in actions
)

types_html = "".join(f'<span class="badge2">{t}</span>' for t in activity_types)

html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>D：天猫运营 · 执行单 {TS.replace("_", " ")}</title>
<style>
  *{{box-sizing:border-box;margin:0;padding:0}}
  body{{font-family:-apple-system,"PingFang SC","Microsoft YaHei",Helvetica,Arial,sans-serif;background:#f4f6fa;color:#1f2937;line-height:1.6;padding:24px}}
  .wrap{{max-width:1120px;margin:0 auto}}
  .hero{{background:linear-gradient(135deg,#1f4e79 0%,#2f7bbf 55%,#3fa9d8 100%);color:#fff;border-radius:16px;padding:28px 30px;box-shadow:0 10px 26px rgba(31,78,121,.28)}}
  .hero h1{{font-size:23px;font-weight:700;letter-spacing:.3px}}
  .hero .sub{{margin-top:8px;font-size:12.5px;opacity:.92;font-family:ui-monospace,Consolas,monospace;word-break:break-all}}
  .badges{{margin-top:16px;display:flex;flex-wrap:wrap;gap:8px}}
  .badge{{background:rgba(255,255,255,.18);border:1px solid rgba(255,255,255,.38);border-radius:999px;padding:4px 12px;font-size:12px}}
  .alert{{margin:18px 0;border-radius:12px;padding:14px 18px;font-size:13.5px;border-left:5px solid}}
  .alert.warn{{background:#fff8e6;border-color:#e8a33d;color:#7a4e07}}
  .alert b{{font-weight:700}}
  .kpis{{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin:20px 0}}
  .kpi{{background:#fff;border-radius:14px;padding:16px 18px;box-shadow:0 3px 14px rgba(23,43,77,.08);border-top:4px solid #2f7bbf}}
  .kpi.k2{{border-top-color:#34a853}}.kpi.k3{{border-top-color:#e8a33d}}.kpi.k4{{border-top-color:#8b5cf6}}
  .kpi .lab{{font-size:12px;color:#6b7280;letter-spacing:.4px}}
  .kpi .val{{font-size:26px;font-weight:700;margin:6px 0 2px;color:#12314f}}
  .kpi .val small{{font-size:13px;font-weight:500;color:#6b7280}}
  .kpi .hint{{font-size:11.5px;color:#9aa3b0}}
  section{{background:#fff;border-radius:14px;padding:20px 22px;margin:18px 0;box-shadow:0 3px 14px rgba(23,43,77,.07)}}
  h2{{font-size:16px;color:#12314f;margin-bottom:14px;padding-left:11px;border-left:4px solid #2f7bbf;line-height:1.2}}
  table{{width:100%;border-collapse:collapse;font-size:13px}}
  th{{background:#eef4fb;color:#12314f;font-weight:600;text-align:left;padding:9px 10px;border-bottom:2px solid #cfe0f2;white-space:nowrap}}
  td{{padding:9px 10px;border-bottom:1px solid #eef1f5;vertical-align:top}}
  td.num{{text-align:center;color:#8892a0;width:34px}}
  td.src{{color:#0f4c81;font-size:12.5px}}
  td.modcell{{background:#f7fbff;font-weight:700;color:#12314f;white-space:nowrap;text-align:center;vertical-align:middle}}
  td.ok{{color:#1e6b38;background:#f4fbf7}}
  td.bad{{color:#8a3030;background:#fdf6f6}}
  tr:last-child td{{border-bottom:none}}
  tr:hover td{{background:#fafcff}}
  code{{font-family:ui-monospace,Consolas,monospace;background:#f1f5f9;border-radius:5px;padding:1px 6px;font-size:12px;color:#0f4c81}}
  .pill{{display:inline-block;border-radius:999px;padding:2px 10px;font-size:11.5px;font-weight:600;white-space:nowrap}}
  .p-todo{{background:#fdf0d5;color:#8a5a04}}
  .p-ok{{background:#e6f5ec;color:#1e6b38}}
  .p-core{{background:#e8effc;color:#2b4c9b}}
  .p-info{{background:#f0eafc;color:#5b33b5}}
  .badge2{{display:inline-block;background:#eef4fb;border:1px solid #cfe0f2;color:#12314f;border-radius:8px;padding:3px 10px;font-size:12.5px;margin:0 6px 6px 0}}
  ul.act{{list-style:none}}
  ul.act li{{position:relative;padding:9px 0 9px 30px;border-bottom:1px dashed #e8edf3;font-size:13.5px}}
  ul.act li:last-child{{border-bottom:none}}
  ul.act li::before{{content:"▸";position:absolute;left:10px;color:#2f7bbf;font-weight:700}}
  ul.act li b{{color:#12314f}}
  ul.act li .who{{color:#8a5a04;background:#fdf6e7;border-radius:5px;padding:1px 7px;font-size:11.5px;margin-left:6px}}
  .note{{font-size:12.5px;color:#6b7280;margin-top:10px;background:#f8fafc;border-radius:10px;padding:12px 14px}}
  footer{{margin:22px 0 40px;font-size:12px;color:#8b95a3;text-align:center;line-height:1.9}}
  footer b{{color:#5b6674}}
  @media(max-width:900px){{.kpis{{grid-template-columns:repeat(2,1fr)}}}}
</style>
</head>
<body>
<div class="wrap">

  <div class="hero">
    <h1>D：天猫运营 · 执行单</h1>
    <div class="sub">房间 room-xgoibqb8（UN项目工程） &nbsp;|&nbsp; 节点 uid {SOFTP} &nbsp;|&nbsp; 执行时间 {TS.replace("_"," ")}</div>
    <div class="badges">
      <span class="badge">路径：C：UN项目利润分 → P：渠道运营 → 天猫 → 周中昱 → D：天猫运营</span>
      <span class="badge">责任人：周中昱</span>
      <span class="badge">子树实测 78 节点 · 4 子模块</span>
      <span class="badge">traceId {TRACE}</span>
      <span class="badge">状态：规则已确认 · 数值待取数</span>
    </div>
  </div>

  <div class="alert warn">
    <b>执行结论：规则侧完整，数值侧阻塞。</b>
    本 SOP「D：天猫运营」下辖 <b>4 个子模块</b>（活动运营 / 货品运营 / 直播运营 / 会员运营），其步骤与检查项已 <b>100% 从导图真实子树取到</b>（78 节点，含 4 项活动检查、达播/店播 ROI 判定、链接与页面验证、会员四项指标口径）。
    但 <b>全部数值指标</b>（达播/店播 ROI、会员入会率与复购率、链接转化率/客单价等）依赖天猫生意参谋、千牛、会员中心及达播结算后台，<b>导图未挂数据源</b>；
    唯一外链「检查文案，符合大贸备案要求」指向企业微信文档（实测 HTTP 200 但为登录壳 120,113 字节、0 命中）。因此本单交付 <b>「规则已确认 + 监控台账（待回填）」</b>，<b>未填入任何推测数字</b>。
  </div>

  <div class="kpis">
    <div class="kpi">
      <div class="lab">覆盖子模块</div>
      <div class="val">4 <small>个</small></div>
      <div class="hint">活动运营 · 货品运营 · 直播运营 · 会员运营</div>
    </div>
    <div class="kpi k2">
      <div class="lab">导图实测节点 / 叶子检查项</div>
      <div class="val">78 <small>节点 / 48 叶子</small></div>
      <div class="hint">实时子树搜索，规则完整</div>
    </div>
    <div class="kpi k3">
      <div class="lab">已取到数值</div>
      <div class="val">0 <small>条</small></div>
      <div class="hint">数值均在店铺后台/结算后台，导图未挂</div>
    </div>
    <div class="kpi k4">
      <div class="lab">已挂数据源外链</div>
      <div class="val">1 <small>个</small></div>
      <div class="hint">大贸备案表（企业微信登录墙）</div>
    </div>
  </div>

  <section>
    <h2>SOP 四大子模块结构（导图原文）</h2>
    <table>
      <thead><tr><th>子模块</th><th>节点 uid（前 8 位）</th><th>子树节点数</th><th>主要内容</th></tr></thead>
      <tbody>
        {mod_rows}
      </tbody>
    </table>
    <div class="note">活动运营适用活动类型：{types_html}</div>
  </section>

  <section>
    <h2>SOP 步骤与检查项（按子模块）</h2>
    <table>
      <thead><tr><th>子模块</th><th>#</th><th>步骤 / 检查项（导图原文）</th><th>类型</th><th>完成标志</th><th>数据源</th><th>状态</th></tr></thead>
      <tbody>
        {rows_html(activity_rows, "活动运营")}
        {rows_html(huopin_rows, "货品运营")}
        {rows_html(zhibo_rows, "直播运营")}
        {rows_html(huiyuan_rows, "会员运营")}
      </tbody>
    </table>
    <div class="note">「完成标志」为对导图检查项的落地判定条件；「数据源」为实际取数位置，导图内除大贸备案表外均未挂外链。</div>
  </section>

  <section>
    <h2>判定口径表（完成 / 未完成分支）</h2>
    <table>
      <thead><tr><th>环节</th><th>指标 / 阈值</th><th>完成分支</th><th>未完成分支</th><th>数据源</th></tr></thead>
      <tbody>
        {criteria_rows}
      </tbody>
    </table>
  </section>

  <section>
    <h2>本周动作</h2>
    <ul class="act">
      {actions_html}
    </ul>
  </section>

  <footer>
    <b>数据来源：</b>导图 room-xgoibqb8 实测子树（节点 uid {SOFTP}，traceId {TRACE}）· 执行时间 {TS.replace("_"," ")}<br>
    <b>缺口说明：</b>① 唯一外链「大贸备案表」为企业微信文档登录墙，匿名不可读；② 达播/店播 ROI、会员与链接类指标数值均在天猫生意参谋 / 千牛 / 会员中心 / 达播结算后台，导图未挂外链；③ 直播运营节点在分页返回中 text 字段为空富文本片段，题名以 content_fragment 还原为「D：直播运营」。<br>
    本单不含任何推测数字，数值留空待回填。
  </footer>

</div>
</body>
</html>
"""

with open(os.path.join(OUT, HTML_NAME), "w", encoding="utf-8") as f:
    f.write(html)
print("HTML ->", os.path.join(OUT, HTML_NAME))

# ---------------------------------------------------------------- XLSX
wb = Workbook()
thin = Side(style="thin", color="D6E0EC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_font = Font(color="FFFFFF", bold=True, size=11)
title_font = Font(color="12314F", bold=True, size=13)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_header(ws, row=1, ncol=6):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center
        cell.border = border

def auto_width(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# Sheet 1 总览
ws = wb.active
ws.title = "总览"
ws["A1"] = "D：天猫运营 · 监控台账（数值待回填）"
ws["A1"].font = Font(bold=True, size=15, color="1F4E79")
ws.merge_cells("A1:D1")
meta = [
    ("SOP", "D：天猫运营"),
    ("节点 uid", SOFTP),
    ("路径", "C：UN项目利润分 → P：渠道运营 → 天猫 → 周中昱 → D：天猫运营"),
    ("责任人", "周中昱"),
    ("traceId", TRACE),
    ("房间", "room-xgoibqb8（UN项目工程）"),
    ("执行时间", TS.replace("_", " ")),
    ("子树节点数", "78（4 子模块）"),
    ("已取到数值", "0 条（数值待取数，导图未挂数据源）"),
    ("外链数据源", "大贸备案表：https://doc.weixin.qq.com/sheet/e3_AbUAJQbFAEECNnLpgSXeQRy0UaDE9（企业微信登录墙）"),
]
r = 3
for k, v in meta:
    ws.cell(row=r, column=1, value=k).font = Font(bold=True, color="12314F")
    ws.cell(row=r, column=1).border = border
    c = ws.cell(row=r, column=2, value=v); c.alignment = wrap; c.border = border
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r += 1
r += 1
ws.cell(row=r, column=1, value="子模块").font = Font(bold=True, color="1F4E79"); r += 1
for h, col in zip(["子模块", "节点uid", "子树节点数", "主要内容"], range(1, 5)):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, 4); r += 1
for n, u, c, d in modules:
    ws.cell(row=r, column=1, value=n).border = border
    ws.cell(row=r, column=2, value=u).border = border
    ws.cell(row=r, column=3, value=c).border = border
    ws.cell(row=r, column=4, value=d).alignment = wrap; ws.cell(row=r, column=4).border = border
    r += 1
auto_width(ws, [22, 40, 12, 46, 20])

def add_rows_sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(headers))
    r = 2
    for row in rows:
        for i, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.alignment = wrap
            cell.border = border
        r += 1
    ws.freeze_panes = "A2"
    auto_width(ws, widths)
    return ws

step_hdr = ["#", "步骤 / 检查项", "类型", "完成标志", "数据源", "状态", "实际值（待填）"]
step_w = [5, 46, 12, 40, 18, 10, 18]

add_rows_sheet("活动运营", ["#", "步骤 / 检查项", "类型", "完成标志", "数据源", "状态", "实际值（待填）"],
               [(i, r[0], r[1], r[2], r[3], r[4], "") for i, r in enumerate(activity_rows, 1)], step_w)
add_rows_sheet("货品运营", ["#", "步骤 / 检查项", "类型", "完成标志", "数据源", "状态", "实际值（待填）"],
               [(i, r[0], r[1], r[2], r[3], r[4], "") for i, r in enumerate(huopin_rows, 1)], step_w)
add_rows_sheet("直播运营", ["#", "步骤 / 检查项", "类型", "完成标志", "数据源", "状态", "实际值（待填）"],
               [(i, r[0], r[1], r[2], r[3], r[4], "") for i, r in enumerate(zhibo_rows, 1)], step_w)
add_rows_sheet("会员运营", ["#", "步骤 / 检查项", "类型", "完成标志", "数据源", "状态", "实际值（待填）"],
               [(i, r[0], r[1], r[2], r[3], r[4], "") for i, r in enumerate(huiyuan_rows, 1)], step_w)

add_rows_sheet("判定口径", ["环节", "指标 / 阈值", "完成分支", "未完成分支", "数据源"],
               criteria, [18, 34, 30, 34, 26])

gap_rows = [
    ("达播 ROI", "达播成果验证", "达播结算后台（导图未挂）", "需要达播结算/后台金额与成本数据"),
    ("定播 ROI", "店播成果验证", "店播中控后台（导图未挂）", "需要店播中控投放与成交数据"),
    ("链接转化率/客单价/总销售额", "链接是否成功验证", "生意参谋（导图未挂）", "需生意参谋链接级数据"),
    ("页面合规 / 文案备案", "优化页面·合规验证", "大贸备案表（企业微信登录墙）", "需已登录企业微信打开备案表"),
    ("新会员入会率 / 成交占比", "新会员规划验证", "天猫会员中心（导图未挂）", "需会员中心导出"),
    ("会员复购率 / 整体销售占比", "会员成交验证", "天猫会员中心（导图未挂）", "需会员中心导出"),
]
add_rows_sheet("取数缺口清单", ["指标", "所属环节", "数据源 / 获取位置", "缺口说明"],
               gap_rows, [28, 22, 34, 40])

xlsx_path = os.path.join(OUT, XLSX_NAME)
wb.save(xlsx_path)
print("XLSX ->", xlsx_path)
